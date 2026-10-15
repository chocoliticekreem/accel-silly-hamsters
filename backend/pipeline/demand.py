"""demand.json: hourly arrivals per map entry point, the returning pool and the club nights.

Units are SIM agents per hour (expected values), not real people: the weekday total is fixed at
SIM_WEEKDAY_ARRIVALS and every real count is scaled by one factor, written into the file's notes.

Measured (TfL open data, no key needed):
  NUMBAT 2025 typical-day workbooks, one per day type: Shoreditch High Street station exits by
  quarter hour. Exits are people arriving into the area, so they give the station's arrivals by
  hour, the weekend/weekday volume ratio and, in the Fri/Sat 00:00-05:00 bands, the Night
  Overground.
  Unified API, Windrush line Night route: confirms the night service calls at the station.
Assumed (named constants below, repeated in the notes): the station's share of all arrivals, the
street-entry timing mix, club fill and the regulars model. Club events are not scraped (Resident
Advisor, Dice and Eventbrite restrict it); club nights come from the map's club opening hours.
"""

from __future__ import annotations

import argparse
import io
import logging
import math
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from typing import Any

import httpx
import openpyxl

from backend import cache, store, timeutil
from backend.contract import Building, EntryPoint, MapData, MapStore, Source, Street
from backend.pipeline import common
from backend.schemas import ClubNight, DemandFile, EntryDemand

log = logging.getLogger(__name__)

STEP = "demand"
# Bump when a code change alters demand.json for the same map, so the pipeline cache rebuilds it.
MODEL_VERSION = 3

# --- TfL NUMBAT: station exits by quarter hour ----------------------------------------------

NUMBAT_URL = "https://crowding.data.tfl.gov.uk/NUMBAT/NUMBAT%202025/NBT25{day_type}_Outputs.xlsx"
NUMBAT_DAY_TYPES = {
    "MON": "Monday",
    "TWT": "Tuesday-Thursday",
    "FRI": "Friday",
    "SAT": "Saturday",
    "SUN": "Sunday",
}
NIGHT_LABELS = {
    "MON": "Mon",
    "TWT": "Tue-Thu (each)",
    "FRI": "Fri",
    "SAT": "Sat",
    "SUN": "Sun",
}
STATION_NAME = "Shoreditch High Street"
STATION_NAPTAN = "910GSHRDHST"
EXITS_SHEET = "Station_Exits"
QUARTER_HOUR_LABEL = re.compile(r"^(\d{2})(\d{2})-\d{4}$")  # "0500-0515"
QUARTER_HOURS_PER_DAY = 96
MAX_DOWNLOAD_BYTES = 20_000_000  # each NUMBAT workbook is ~12 MB; refuse anything far bigger

# NUMBAT traffic days run 05:00-05:00, so a workbook's 00:00-04:59 bands are the next calendar
# morning. Calendar day -> (workbook for its 05:00-23:59, workbook whose 00:00-04:59 it inherits).
TRAFFIC_DAY_START_HOUR = 5
CALENDAR_SOURCES = {
    "mon": ("MON", "SUN"),
    "tue": ("TWT", "MON"),
    "wed": ("TWT", "TWT"),
    "thu": ("TWT", "TWT"),
    "fri": ("FRI", "TWT"),
    "sat": ("SAT", "FRI"),
    "sun": ("SUN", "SAT"),
}
WEEKEND_DAYS = ("sat", "sun")
WORKING_DAYS = tuple(d for d in timeutil.WEEKDAYS if d not in WEEKEND_DAYS)

# Used only when NUMBAT is unreachable: calendar_profiles() of the five 2025 workbooks as fetched
# on SNAPSHOT_DATE (mean SHS exits per clock hour). The file is then flagged fallback=True.
SNAPSHOT_DATE = "2026-10-15"
NUMBAT_FOLDER_URL = "https://crowding.data.tfl.gov.uk/#!NUMBAT%2FNUMBAT%202025%2F"
# fmt: off
SNAPSHOT_EXITS = {
    "weekday": [
        24.6, 0.0, 0.0, 0.0, 0.0, 8.9, 201.4, 733.4, 1807.2, 1802.2, 911.4, 600.0,
        560.2, 497.6, 440.1, 453.0, 538.2, 693.9, 799.1, 576.1, 335.6, 209.6, 178.5, 123.9,
    ],
    "weekend": [
        145.5, 47.4, 22.9, 13.8, 7.8, 9.9, 49.8, 97.1, 223.9, 440.8, 662.0, 838.3,
        957.7, 990.3, 911.1, 826.2, 735.9, 680.2, 639.2, 579.4, 429.9, 296.5, 238.3, 179.0,
    ],
}
# fmt: on

NIGHT_ROUTE_URL = "https://api.tfl.gov.uk/Line/windrush/Route/Sequence/outbound"

# --- scale and entry points -----------------------------------------------------------------

SIM_WEEKDAY_ARRIVALS = 800.0  # agents per weekday via all entry points (BUILD_NOTES: 600-1200)
# Assumption: share of all arrivals into the area that come through Shoreditch High Street. The
# rest walk in from Liverpool Street / Old Street, buses, bikes and neighbouring streets.
STATION_SHARE = 0.30
# Street arrivals are not measured by hour here (TfL BUSTO bus files are 55-125 MB). They reuse
# the station's exit timing plus this share of night-time-economy walk-ins.
STREET_LEISURE_SHARE = {"weekday": 0.30, "weekend": 0.40}
# Relative weight of those walk-ins by clock hour: (start hour, end hour, weight).
LEISURE_BANDS: dict[str, list[tuple[int, int, float]]] = {
    "weekday": [(0, 1, 0.5), (11, 14, 2.0), (14, 17, 1.0), (17, 23, 3.0), (23, 24, 1.5)],
    "weekend": [(0, 2, 2.0), (2, 4, 1.0), (10, 12, 1.0), (12, 20, 2.0), (20, 24, 3.0)],
}
STREET_SNAP_M = 40.0  # an entry within this distance of a street takes that street's width
DEFAULT_STREET_WIDTH_M = 8.0

# --- clubs ----------------------------------------------------------------------------------

CLUB_DAYS = ["thu", "fri", "sat"]
CLUB_DEFAULT_HOURS = ("22:00", "03:00")
# Map hours count as a club night only when they open from 18:00 and close after midnight, by
# 08:00. OSM hours like 12:00-00:00 describe a venue's bar trade, and the sim takes club hours
# (and so the spill-out time) from ClubNight.
CLUB_NIGHT_OPENS_FROM = "18:00"
CLUB_NIGHT_CLOSES_BY = "08:00"
CLUB_M2_PER_PERSON = 0.5  # Approved Document B vol 2, Table D1: "club, dance floor or hall"
CLUB_PEAK_FILL = 0.6  # assumption: typical peak occupancy vs the fire-safety maximum
CLUB_REAL_CAPACITY_RANGE = (150.0, 1500.0)
CLUB_DEFAULT_REAL_CAPACITY = 500.0  # assumption when the map has no floor area for the club
ADB_SOURCE = common.cite(
    "Approved Document B vol 2 (2026), Table D1: floor space factor 0.5 m2/person for clubs "
    "(static reference, checked 2026-10-15)",
    "https://assets.publishing.service.gov.uk/media/6aa3bb7492e72b8ac437f145/AD_B2_2026.pdf",
)

# --- returning pool -------------------------------------------------------------------------

AM_PEAK_HOURS = range(7, 10)  # TfL's AM peak band (07:00-10:00): commuters, regular by nature
REGULAR_SHARE_OF_OTHERS = 0.25  # assumption: residents, regular diners/drinkers, club regulars
REGULAR_ATTENDANCE = 0.9  # assumption: a regular arrives on ~9 of 10 days of the busier day type

CHECKED_NOT_USED = (
    "Checked 2026-10-15, not used: TfL crowding API (/crowding/910GSHRDHST returns "
    "isFound=false); ONS Time Use Survey tables (OTUS, March 2024) give daily minutes per "
    "activity with no hour-of-day split, so hour-of-day comes from NUMBAT; TfL BUSTO bus demand "
    "files are 55-125 MB (over the download cap); Resident Advisor, Dice and Eventbrite are not "
    "scraped (terms)."
)


class SourceFormatError(ValueError):
    """A downloaded source did not have the expected layout."""


@dataclass
class StationExits:
    """Mean SHS exits (real people) per clock hour on a weekday and on a weekend calendar day."""

    weekday: list[float]
    weekend: list[float]
    sources: list[Source]
    measured: bool
    night_exits: dict[str, float] = field(default_factory=dict)  # NUMBAT day type -> 01:00-05:00
    error: str | None = None


# --- build ----------------------------------------------------------------------------------


def build(map_data: MapData | None = None, refresh: bool = False) -> DemandFile:
    """Build data/processed/demand.json for `map_data` (default: data/processed/map.json).

    Rebuilds only when the map's entry points, streets or clubs (or MODEL_VERSION) changed, or
    when the previous build fell back to the snapshot. refresh=True re-downloads the TfL sources.
    """
    if map_data is None:
        map_data = store.load_map()
    inputs = cache_inputs(map_data)
    if not refresh and cache.pipeline_up_to_date(STEP, inputs):
        previous = _load_previous()
        if previous is not None and not previous.fallback:
            return previous
    prior_refresh = cache.policy.refresh
    cache.set_policy(refresh=prior_refresh or refresh)
    try:
        demand = build_demand(map_data)
    finally:
        cache.set_policy(refresh=prior_refresh)
    common.write_processed("demand.json", demand)
    cache.pipeline_mark(STEP, inputs)
    return demand


def cache_inputs(map_data: MapData) -> dict[str, Any]:
    """Everything that changes demand.json: model version, sources, entries, streets, clubs."""
    clubs = _clubs(map_data)
    club_buildings = {c.buildingId for c in clubs}
    return {
        "model": MODEL_VERSION,
        "numbat": [numbat_url(d) for d in NUMBAT_DAY_TYPES],
        "entryPoints": [e.model_dump(mode="json") for e in map_data.entryPoints],
        "streets": [s.model_dump(mode="json") for s in map_data.streets],
        "clubs": [c.model_dump(mode="json") for c in clubs],
        "clubBuildings": [
            b.model_dump(mode="json") for b in map_data.buildings if b.id in club_buildings
        ],
    }


def build_demand(map_data: MapData) -> DemandFile:
    """Compute demand.json for a map. No caching of the result and no file writes."""
    if not map_data.entryPoints:
        raise ValueError("map.json has no entry points; demand needs at least one")
    exits = station_exit_profiles()
    night_source, night_calls_station = night_route_check()

    rail_weekday, rail_weekend = sum(exits.weekday), sum(exits.weekend)
    real_weekday = rail_weekday / STATION_SHARE
    agents_per_person = SIM_WEEKDAY_ARRIVALS / real_weekday
    totals = {
        "weekday": SIM_WEEKDAY_ARRIVALS,
        "weekend": SIM_WEEKDAY_ARRIVALS * rail_weekend / rail_weekday,
    }
    shapes = {"weekday": _normalise(exits.weekday), "weekend": _normalise(exits.weekend)}
    entries, street_widths = allocate_entries(map_data, shapes, totals)
    nights = club_nights(map_data, agents_per_person)
    share, pool, am_shares = returning_estimate(entries, totals)

    sources = [*exits.sources, *([night_source] if night_source else []), ADB_SOURCE]
    notes = [
        f"Units: sim agents per hour (expected values; sample e.g. Poisson per hour), not real "
        f"people. Scale: 1 sim agent ~ {1 / agents_per_person:.0f} real arrivals.",
        f"Scale derivation: SHS station exits {rail_weekday:,.0f} per weekday and "
        f"{rail_weekend:,.0f} per weekend day ({_numbat_label(exits)}); assumed station share "
        f"{STATION_SHARE:.0%} of all arrivals into the area gives {real_weekday:,.0f} real "
        f"arrivals per weekday; weekday sim total fixed at {SIM_WEEKDAY_ARRIVALS:.0f} (BUILD_NOTES "
        f"budget 600-1200/day); weekend total {totals['weekend']:.0f} follows the measured "
        f"weekend/weekday exit ratio {rail_weekend / rail_weekday:.2f}.",
        "Station entries: hour-of-day = SHS exits (people arriving into the area; quarter-hour "
        "bands summed to hours), TfL traffic days (05:00-05:00) mapped to calendar days; several "
        "station entries share the station budget equally.",
        _night_note(exits, night_calls_station),
        _street_note(street_widths),
        _club_note(map_data, nights, agents_per_person),
        f"returningShare and returningPoolSize (estimate; open question in the spec): commuters "
        f"= arrivals in the TfL AM peak band 07:00-10:00 ({am_shares['weekday']:.0%} of weekday, "
        f"{am_shares['weekend']:.0%} of weekend arrivals) plus {REGULAR_SHARE_OF_OTHERS:.0%} of "
        f"all other arrivals assumed regulars (residents, regular diners and drinkers, club "
        f"regulars); weekly mix = {share:.2f}. Pool = share x busiest daily total / "
        f"{REGULAR_ATTENDANCE} (a regular arrives on ~9 of 10 days) = {pool}.",
        CHECKED_NOT_USED,
    ]
    if not exits.measured:
        notes.insert(
            0,
            f"FALLBACK: TfL NUMBAT unavailable ({exits.error}); station timing and volumes use the "
            f"snapshot of the 2025 workbooks embedded in demand.py (taken {SNAPSHOT_DATE}).",
        )
    return DemandFile(
        entryPoints=entries,
        returningPoolSize=pool,
        returningShare=share,
        clubNights=nights,
        sources=sources,
        fallback=not exits.measured,
        notes=notes,
    )


def daily_totals(demand: DemandFile) -> tuple[float, float]:
    """(weekday, weekend) arrivals per day summed over every entry point, clubs excluded."""
    weekday = sum(sum(e.weekday) for e in demand.entryPoints.values())
    weekend = sum(sum(e.weekend) for e in demand.entryPoints.values())
    return weekday, weekend


# --- station profile (TfL NUMBAT) -----------------------------------------------------------


def numbat_url(day_type: str) -> str:
    return NUMBAT_URL.format(day_type=day_type)


def station_exit_profiles() -> StationExits:
    """SHS exits per hour from the five NUMBAT workbooks, or the embedded snapshot on failure."""
    by_workbook: dict[str, list[float]] = {}
    sources: list[Source] = []
    try:
        for day_type, label in NUMBAT_DAY_TYPES.items():
            hourly, fetched_at = numbat_hourly_exits(day_type)
            by_workbook[day_type] = hourly
            sources.append(
                common.cite(
                    f"TfL NUMBAT 2025, typical {label}: {STATION_NAME} station exits by quarter "
                    "hour",
                    numbat_url(day_type),
                    fetched_at,
                )
            )
        profiles = calendar_profiles(by_workbook)
        if min(sum(profiles["weekday"]), sum(profiles["weekend"])) <= 0:
            raise SourceFormatError(f"{STATION_NAME} has no exits in the NUMBAT workbooks")
    except (common.FetchError, SourceFormatError) as e:
        log.warning("demand: TfL NUMBAT unavailable, using the embedded snapshot: %s", e)
        return StationExits(
            weekday=list(SNAPSHOT_EXITS["weekday"]),
            weekend=list(SNAPSHOT_EXITS["weekend"]),
            sources=[
                common.cite(
                    f"TfL NUMBAT 2025 station exits, snapshot embedded in demand.py (taken "
                    f"{SNAPSHOT_DATE}; live source unavailable, fallback)",
                    NUMBAT_FOLDER_URL,
                )
            ],
            measured=False,
            error=str(e),
        )
    night = {d: sum(h[1:TRAFFIC_DAY_START_HOUR]) for d, h in by_workbook.items()}
    return StationExits(profiles["weekday"], profiles["weekend"], sources, True, night)


def numbat_hourly_exits(day_type: str) -> tuple[list[float], str]:
    """SHS exits per clock hour over one NUMBAT traffic day, and when the workbook was fetched.

    The scrape cache (data/raw/tfl/) keeps the extracted quarter-hour row, not the ~12 MB file.
    """
    url = numbat_url(day_type)
    key = {"extract": EXITS_SHEET, "station": STATION_NAME}
    cached = cache.scrape_get("tfl", url, key)
    if cached is None:
        fetched_at = store.now_iso()
        cached = {
            "content": parse_station_exits(_download(url), STATION_NAME),
            "fetchedAt": fetched_at,
        }
        cache.scrape_put("tfl", url, key, cached, fetched_at)
    return hourly_from_quarter_hours(cached["content"]), cached["fetchedAt"]


def _download(url: str) -> bytes:
    """GET a binary file (common.fetch caches text/JSON only). Raises common.FetchError."""
    try:
        with httpx.stream(
            "GET",
            url,
            headers={"User-Agent": common.USER_AGENT},
            timeout=common.TIMEOUT,
            follow_redirects=True,
        ) as response:
            response.raise_for_status()
            chunks: list[bytes] = []
            size = 0
            for chunk in response.iter_bytes():
                size += len(chunk)
                if size > MAX_DOWNLOAD_BYTES:
                    raise common.FetchError(f"tfl: {url} is over {MAX_DOWNLOAD_BYTES} bytes")
                chunks.append(chunk)
    except httpx.HTTPError as e:
        raise common.FetchError(f"tfl: GET {url} failed: {e}") from e
    return b"".join(chunks)


def parse_station_exits(xlsx: bytes, station: str) -> dict[str, float]:
    """Quarter-hour exits for `station` from a NUMBAT workbook, keyed by band start "HHMM"."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError) as e:
        raise SourceFormatError(f"NUMBAT workbook unreadable: {e}") from e
    try:
        if EXITS_SHEET not in workbook.sheetnames:
            raise SourceFormatError(f"NUMBAT workbook has no {EXITS_SHEET} sheet")
        header: tuple[Any, ...] | None = None
        station_col = -1
        for row in workbook[EXITS_SHEET].iter_rows(values_only=True):
            if header is None:
                if "Station" in row and any(_quarter_hour_start(c) for c in row):
                    header, station_col = row, row.index("Station")
                continue
            if station_col < len(row) and row[station_col] == station:
                return _quarter_hour_values(header, row)
        raise SourceFormatError(f"{station!r} not found in the {EXITS_SHEET} sheet")
    finally:
        workbook.close()


def _quarter_hour_start(label: Any) -> str | None:
    match = QUARTER_HOUR_LABEL.match(label) if isinstance(label, str) else None
    return f"{match.group(1)}{match.group(2)}" if match else None


def _quarter_hour_values(header: tuple[Any, ...], row: tuple[Any, ...]) -> dict[str, float]:
    values: dict[str, float] = {}
    for i, label in enumerate(header):
        start = _quarter_hour_start(label)
        if start is None:
            continue
        cell = row[i] if i < len(row) else None
        if cell is not None and not isinstance(cell, int | float):
            raise SourceFormatError(f"non-numeric exits {cell!r} in band {label}")
        values[start] = float(cell or 0.0)
    if len(values) != QUARTER_HOURS_PER_DAY:
        raise SourceFormatError(
            f"expected {QUARTER_HOURS_PER_DAY} quarter hours, got {len(values)}"
        )
    return values


def hourly_from_quarter_hours(quarter_hours: dict[str, float]) -> list[float]:
    """Sum "HHMM"-keyed quarter-hour counts into 24 clock-hour totals."""
    hourly = [0.0] * 24
    for start, value in quarter_hours.items():
        hourly[int(start[:2]) % 24] += value
    return hourly


def calendar_profiles(by_workbook: dict[str, list[float]]) -> dict[str, list[float]]:
    """Mean exits per clock hour on a weekday and on a weekend calendar day.

    Each calendar day takes 05:00-23:59 from its own workbook and 00:00-04:59 from the previous
    traffic day's (CALENDAR_SOURCES), so Fri/Sat-night trains land on Saturday/Sunday mornings.
    """
    days = {
        day: [
            by_workbook[previous][h] if h < TRAFFIC_DAY_START_HOUR else by_workbook[own][h]
            for h in range(24)
        ]
        for day, (own, previous) in CALENDAR_SOURCES.items()
    }
    return {
        "weekday": _mean_by_hour([days[d] for d in WORKING_DAYS]),
        "weekend": _mean_by_hour([days[d] for d in WEEKEND_DAYS]),
    }


def night_route_check() -> tuple[Source | None, bool | None]:
    """Whether TfL lists SHS on the Windrush line's Night route; (None, None) when unreachable."""
    try:
        body, fetched_at = common.fetch("tfl", NIGHT_ROUTE_URL, params={"serviceTypes": "Night"})
    except (common.FetchError, ValueError) as e:
        log.warning("demand: Night route check failed: %s", e)
        return None, None
    if not isinstance(body, dict):
        return None, None
    stops = [sp for seq in body.get("stopPointSequences", []) for sp in seq.get("stopPoint", [])]
    calls = any(
        STATION_NAPTAN in (sp.get("id"), sp.get("stationId"), sp.get("topMostParentId"))
        for sp in stops
    )
    source = common.cite(
        "TfL Unified API: Windrush line Night route stop sequence",
        f"{NIGHT_ROUTE_URL}?serviceTypes=Night",
        fetched_at,
    )
    return source, calls


# --- entry points ---------------------------------------------------------------------------


def leisure_shape(day_type: str) -> list[float]:
    """Night-time-economy walk-in arrivals by clock hour (assumption), normalised to sum 1."""
    weights = [0.0] * 24
    for start, end, weight in LEISURE_BANDS[day_type]:
        for hour in range(start, end):
            weights[hour] = weight
    return _normalise(weights)


def street_shape(rail_shape: Sequence[float], day_type: str) -> list[float]:
    """Street-entry arrivals by hour: station exit timing mixed with the leisure walk-ins."""
    mix = STREET_LEISURE_SHARE[day_type]
    return [
        (1 - mix) * rail + mix * leisure
        for rail, leisure in zip(rail_shape, leisure_shape(day_type), strict=True)
    ]


def street_width_at(entry: EntryPoint, streets: Sequence[Street]) -> float:
    """Width of the street an entry point sits on (the nearest one within STREET_SNAP_M)."""
    nearest: tuple[float, float] | None = None  # (distance, width)
    for street in streets:
        points = street.points
        segments = list(zip(points, points[1:], strict=False)) or [(points[0], points[0])]
        for a, b in segments:
            distance = _distance_to_segment((entry.x, entry.z), a, b)
            if nearest is None or distance < nearest[0]:
                nearest = (distance, street.width)
    if nearest is None or nearest[0] > STREET_SNAP_M or nearest[1] <= 0:
        return DEFAULT_STREET_WIDTH_M
    return nearest[1]


def allocate_entries(
    map_data: MapData, shapes: dict[str, list[float]], totals: dict[str, float]
) -> tuple[dict[str, EntryDemand], dict[str, float]]:
    """Hourly arrivals for every entry point, and the street width that weighted each street entry.

    Station entries share STATION_SHARE of the daily total and follow the station exit shape;
    street entries share the rest in proportion to the width of the street they sit on.
    """
    stations = [e for e in map_data.entryPoints if e.kind == "station"]
    streets = [e for e in map_data.entryPoints if e.kind == "street"]
    station_share = STATION_SHARE if stations and streets else (1.0 if stations else 0.0)
    widths = {e.id: street_width_at(e, map_data.streets) for e in streets}
    total_width = sum(widths.values())
    hourly: dict[str, dict[str, list[float]]] = {e.id: {} for e in map_data.entryPoints}
    for day_type, total in totals.items():
        station_total = total * station_share
        for entry in stations:
            per_entry = station_total / len(stations)
            hourly[entry.id][day_type] = [per_entry * s for s in shapes[day_type]]
        street_hours = street_shape(shapes[day_type], day_type)
        for entry in streets:
            per_entry = (total - station_total) * widths[entry.id] / total_width
            hourly[entry.id][day_type] = [per_entry * s for s in street_hours]
    entries = {
        entry_id: EntryDemand(weekday=_rounded(h["weekday"]), weekend=_rounded(h["weekend"]))
        for entry_id, h in hourly.items()
    }
    return entries, widths


# --- clubs ----------------------------------------------------------------------------------


def club_night_window(club: MapStore) -> tuple[str, str] | None:
    """The club's map hours when they describe a night (see CLUB_NIGHT_OPENS_FROM), else None."""
    hours = club.defaultHours
    if hours is None:
        return None
    opens, closes = timeutil.to_minutes(hours.open), timeutil.to_minutes(hours.close)
    opens_from = timeutil.to_minutes(CLUB_NIGHT_OPENS_FROM)
    closes_by = timeutil.to_minutes(CLUB_NIGHT_CLOSES_BY)
    if opens >= opens_from and 0 < closes <= closes_by:
        return hours.open, hours.close
    return None


def club_hours(club: MapStore) -> tuple[str, str]:
    """The club's map hours when they describe a night, else CLUB_DEFAULT_HOURS."""
    return club_night_window(club) or CLUB_DEFAULT_HOURS


def club_real_capacity(club: MapStore, building: Building | None) -> float:
    """Typical peak occupancy in real people: floor area / 0.5 m2 per person x CLUB_PEAK_FILL."""
    area = club.floorArea or (building.w * building.d if building else None)
    if not area or area <= 0:
        return CLUB_DEFAULT_REAL_CAPACITY
    low, high = CLUB_REAL_CAPACITY_RANGE
    return min(high, max(low, area / CLUB_M2_PER_PERSON * CLUB_PEAK_FILL))


def club_nights(map_data: MapData, agents_per_person: float) -> list[ClubNight]:
    """One ClubNight per club store; capacity is in sim agents (real occupancy x scale)."""
    buildings = {b.id: b for b in map_data.buildings}
    nights = []
    for club in _clubs(map_data):
        open_, close = club_hours(club)
        real = club_real_capacity(club, buildings.get(club.buildingId))
        nights.append(
            ClubNight(
                clubId=club.id,
                days=list(CLUB_DAYS),
                open=open_,
                close=close,
                capacity=max(1, round(real * agents_per_person)),
            )
        )
    return nights


# --- returning pool -------------------------------------------------------------------------


def returning_estimate(
    entries: dict[str, EntryDemand], totals: dict[str, float]
) -> tuple[float, int, dict[str, float]]:
    """(returningShare, returningPoolSize, AM-peak share per day type) from the arrival shape.

    Regular share of a day = AM-peak share (commuters) + REGULAR_SHARE_OF_OTHERS of the rest;
    the weekly mix weights 5 weekdays and 2 weekend days by their totals. The pool is sized so
    the busiest day's regulars are REGULAR_ATTENDANCE of it.
    """
    days = {"weekday": len(WORKING_DAYS), "weekend": len(WEEKEND_DAYS)}
    am_shares: dict[str, float] = {}
    regulars = 0.0
    for day_type in days:
        by_hour = [sum(getattr(e, day_type)[h] for e in entries.values()) for h in range(24)]
        am_shares[day_type] = sum(by_hour[h] for h in AM_PEAK_HOURS) / sum(by_hour)
        regular_share = am_shares[day_type] + REGULAR_SHARE_OF_OTHERS * (1 - am_shares[day_type])
        regulars += days[day_type] * totals[day_type] * regular_share
    share = regulars / sum(days[d] * totals[d] for d in days)
    pool = math.ceil(share * max(totals.values()) / REGULAR_ATTENDANCE)
    return round(share, 3), pool, am_shares


# --- notes ----------------------------------------------------------------------------------


def _numbat_label(exits: StationExits) -> str:
    if exits.measured:
        return "TfL NUMBAT 2025, mean of calendar days"
    return f"TfL NUMBAT 2025 snapshot of {SNAPSHOT_DATE}, fallback"


def _night_note(exits: StationExits, calls_station: bool | None) -> str:
    if calls_station is None:
        route = "the TfL Night route was not checked this build (API unreachable)"
    elif calls_station:
        route = "TfL lists Shoreditch High Street on the Windrush Night route (Unified API)"
    else:
        route = "TfL's Windrush Night route did NOT list Shoreditch High Street this build"
    if exits.night_exits:
        measured = ", ".join(f"{NIGHT_LABELS[d]} {v:.0f}" for d, v in exits.night_exits.items())
        data = f"NUMBAT SHS exits 01:00-05:00 by night: {measured}"
    else:
        data = "per-night NUMBAT exits not available (snapshot)"
    return (
        f"Night Overground: {route}; {data}. Late station arrivals therefore sit in the 00-04 "
        "hours of the weekend profile (Sat/Sun mornings) and are ~0 on weekday nights."
    )


def _street_note(widths: dict[str, float]) -> str:
    weekday, weekend = STREET_LEISURE_SHARE["weekday"], STREET_LEISURE_SHARE["weekend"]
    split = ", ".join(f"{entry_id} {width:g} m" for entry_id, width in widths.items()) or "none"
    return (
        f"Street entries (assumption): hourly arrivals at street edges are not measured here; "
        f"timing = {1 - weekday:.0%} weekday / {1 - weekend:.0%} weekend SHS exit timing plus "
        f"{weekday:.0%} / {weekend:.0%} night-time-economy walk-ins (weekday 11:00-01:00, peak "
        f"17-23; weekend 10:00-04:00, peak 20-24). Split across street entries by the width of "
        f"the street each sits on: {split}."
    )


def _club_note(map_data: MapData, nights: list[ClubNight], agents_per_person: float) -> str:
    buildings = {b.id: b for b in map_data.buildings}
    clubs = {c.id: c for c in _clubs(map_data)}
    details = []
    for night in nights:
        club = clubs[night.clubId]
        real = club_real_capacity(club, buildings.get(club.buildingId))
        hours = _club_hours_origin(club)
        details.append(
            f"{night.clubId} {night.open}-{night.close} ({hours}), {real:.0f} real -> "
            f"{night.capacity} agents"
        )
    low, high = CLUB_REAL_CAPACITY_RANGE
    default_open, default_close = CLUB_DEFAULT_HOURS
    return (
        f"Club nights: one per club store, nights {'/'.join(CLUB_DAYS)}, hours from map.json "
        f"defaultHours when they describe a night (open from {CLUB_NIGHT_OPENS_FROM}, close after "
        f"midnight by {CLUB_NIGHT_CLOSES_BY}), else {default_open}-{default_close} (assumption; "
        f"club events are not scraped). Club-goers are NOT in entryPoints: "
        f"the sim adds `capacity` club-goers per club night; capacity is already in sim agents "
        f"(do not rescale). capacity = floor area / "
        f"{CLUB_M2_PER_PERSON} m2 per person (Approved Document B) x {CLUB_PEAK_FILL:.0%} typical "
        f"fill, clamped {low:.0f}-{high:.0f} real ({CLUB_DEFAULT_REAL_CAPACITY:.0f} without floor "
        f"area), x scale {agents_per_person:.4f}: {'; '.join(details) or 'no clubs in the map'}."
    )


# --- helpers --------------------------------------------------------------------------------


def _clubs(map_data: MapData) -> list[MapStore]:
    return [s for s in map_data.stores if s.type == "club"]


def _club_hours_origin(club: MapStore) -> str:
    hours = club.defaultHours
    if club_night_window(club):
        return "map hours"
    if hours is None:
        return "default: no map hours"
    return f"default: map hours {hours.open}-{hours.close} are not a night"


def _load_previous() -> DemandFile | None:
    try:
        return store.load_demand()
    except (FileNotFoundError, ValueError):
        return None


def _normalise(values: Sequence[float]) -> list[float]:
    total = sum(values)
    if total <= 0:
        raise ValueError("cannot normalise a profile that sums to zero")
    return [v / total for v in values]


def _mean_by_hour(days: Sequence[Sequence[float]]) -> list[float]:
    return [sum(day[h] for day in days) / len(days) for h in range(24)]


def _rounded(values: Sequence[float]) -> list[float]:
    return [round(v, 2) for v in values]


def _distance_to_segment(
    point: tuple[float, float], a: tuple[float, float], b: tuple[float, float]
) -> float:
    (px, pz), (ax, az), (bx, bz) = point, a, b
    dx, dz = bx - ax, bz - az
    length_sq = dx * dx + dz * dz
    t = 0.0 if length_sq == 0 else ((px - ax) * dx + (pz - az) * dz) / length_sq
    t = max(0.0, min(1.0, t))
    return math.hypot(px - (ax + t * dx), pz - (az + t * dz))


def main(argv: Sequence[str] | None = None) -> None:
    """CLI: python -m backend.pipeline.demand [--refresh]"""
    parser = argparse.ArgumentParser(description="Build data/processed/demand.json")
    parser.add_argument("--refresh", action="store_true", help="re-download the TfL sources")
    args = parser.parse_args(argv)
    demand = build(refresh=args.refresh)
    weekday, weekend = daily_totals(demand)
    print(
        f"demand.json: {len(demand.entryPoints)} entry points, {weekday:.0f} weekday / "
        f"{weekend:.0f} weekend arrivals per day, returning pool {demand.returningPoolSize} "
        f"(share {demand.returningShare}), {len(demand.clubNights)} club nights, "
        f"fallback={demand.fallback}"
    )


if __name__ == "__main__":
    main()
